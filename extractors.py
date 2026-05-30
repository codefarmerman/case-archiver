"""
extractors.py — 多格式文本提取
统一入口 extract_sample(path)，按扩展名分派到对应提取器。
支持：pdf / docx / txt·md / xlsx / ofd / doc(可选) / 图片OCR(可选)。
所有提取器优雅降级：依赖缺失或解析失败时返回空串并记日志，绝不抛出。

可选依赖：
  - .doc   ：pywin32（仅 Windows，需安装 Word）  pip install pywin32
  - 图片OCR：pytesseract + Tesseract-OCR 引擎       pip install pytesseract
"""
from __future__ import annotations

import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from logger import get_logger

log = get_logger()

# 各格式默认采样上限
DEFAULT_CHARS = 1500
DEFAULT_PAGES = 2

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".tiff"}


def extract_sample(file_path: Path, n_chars: int = DEFAULT_CHARS, max_pages: int = DEFAULT_PAGES) -> str:
    """按扩展名分派提取文本前若干字。任何失败返回空串。"""
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _extract_pdf(file_path, n_chars, max_pages)
        if suffix == ".docx":
            return _extract_docx(file_path, n_chars)
        if suffix in (".txt", ".md"):
            return file_path.read_text(encoding="utf-8", errors="ignore")[:n_chars]
        if suffix == ".xlsx":
            return _extract_xlsx(file_path, n_chars)
        if suffix == ".ofd":
            return _extract_ofd(file_path, n_chars)
        if suffix == ".doc":
            return _extract_doc(file_path, n_chars)
        if suffix in IMAGE_EXTS:
            return _extract_image_ocr(file_path, n_chars)
        return ""
    except Exception as e:
        log.warning("提取文本失败 %s: %s", file_path.name, e)
        return ""


# ---------------- 各格式实现 ----------------

def _extract_pdf(file_path: Path, n_chars: int, max_pages: int) -> str:
    import pdfplumber
    with pdfplumber.open(str(file_path)) as pdf:
        text = ""
        for page in pdf.pages[:max_pages]:
            text += (page.extract_text() or "") + "\n"
            if len(text) >= n_chars:
                break
        return text[:n_chars]


def _extract_docx(file_path: Path, n_chars: int) -> str:
    from docx import Document
    doc = Document(str(file_path))
    return "\n".join(p.text for p in doc.paragraphs[:40])[:n_chars]


def _extract_xlsx(file_path: Path, n_chars: int) -> str:
    """读取所有工作表前若干行单元格文本。"""
    import openpyxl
    wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            parts.append(f"[表: {ws.title}]")
            for row in ws.iter_rows(max_row=30, values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
                if sum(len(p) for p in parts) >= n_chars:
                    break
            if sum(len(p) for p in parts) >= n_chars:
                break
    finally:
        wb.close()
    return "\n".join(parts)[:n_chars]


def _extract_ofd(file_path: Path, n_chars: int) -> str:
    """OFD 是 zip 容器，正文在 Doc_N/Pages/.../Content.xml 的 TextCode 节点。
    提取所有 TextCode 文本，无第三方依赖。"""
    texts: list[str] = []
    with zipfile.ZipFile(str(file_path)) as zf:
        content_files = [n for n in zf.namelist() if n.lower().endswith(".xml") and "content" in n.lower()]
        # 兜底：若没有明显 Content.xml，扫描所有 xml
        if not content_files:
            content_files = [n for n in zf.namelist() if n.lower().endswith(".xml")]
        for name in content_files:
            try:
                data = zf.read(name)
                root = ET.fromstring(data)
            except Exception:
                continue
            for elem in root.iter():
                tag = elem.tag.split("}")[-1]  # 去命名空间
                if tag == "TextCode" and elem.text and elem.text.strip():
                    texts.append(elem.text.strip())
            if sum(len(t) for t in texts) >= n_chars:
                break
    return "".join(texts)[:n_chars]


def _extract_doc(file_path: Path, n_chars: int) -> str:
    """旧版 .doc 用 Word COM 自动化提取（仅 Windows + 已装 Word）。"""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        log.info(".doc 提取需要 pywin32（pip install pywin32）且安装 Word，已跳过：%s", file_path.name)
        return ""
    word = None
    try:
        import pythoncom  # type: ignore
        pythoncom.CoInitialize()
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(file_path), ReadOnly=True)
        text = doc.Content.Text or ""
        doc.Close(False)
        return text[:n_chars]
    except Exception as e:
        log.warning(".doc 提取失败 %s: %s", file_path.name, e)
        return ""
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass


def _extract_image_ocr(file_path: Path, n_chars: int) -> str:
    """图片 OCR（需 pytesseract + Tesseract 引擎，含中文语言包）。"""
    try:
        import pytesseract  # type: ignore
        from PIL import Image
    except ImportError:
        log.info("图片 OCR 需要 pytesseract（pip install pytesseract）+ Tesseract 引擎，已跳过：%s", file_path.name)
        return ""
    try:
        img = Image.open(str(file_path))
        text = pytesseract.image_to_string(img, lang="chi_sim+eng")
        return (text or "").strip()[:n_chars]
    except Exception as e:
        log.warning("图片 OCR 失败 %s: %s", file_path.name, e)
        return ""
